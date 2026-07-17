#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Batch probe for temporal credit structure on historical films.

Independent experiment; does not touch the production pipeline.

Run:
  E:\MITAS\venvs\ocr\Scripts\python.exe -B outputs\credit_structure_batch20_probe.py
"""
from __future__ import annotations

import argparse
import importlib.util
import json
import re
import sys
import time
from pathlib import Path
from typing import Any


ROOT = Path(r"E:\MITAS")
DATABASE = ROOT / "Database"
TEMPORAL_PROBE = ROOT / "outputs" / "temporal_credit_structure_probe.py"
DEFAULT_OUT = ROOT / "outputs" / "credit_structure_batch20"


def load_module(name: str, path: Path) -> Any:
    spec = importlib.util.spec_from_file_location(name, str(path))
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    assert spec.loader is not None
    spec.loader.exec_module(mod)
    return mod


def parse_year(name: str) -> int | None:
    m = re.search(r"\b(19|20)\d{2}\b", name)
    return int(m.group(0)) if m else None


def load_clip_meta(folder: Path) -> dict[str, Any]:
    path = folder / "clip.json"
    if not path.exists():
        return {}
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return {}


def discover_cases(limit: int, max_year: int) -> list[dict[str, Any]]:
    cases: list[dict[str, Any]] = []
    for folder in sorted([p for p in DATABASE.iterdir() if p.is_dir()], key=lambda p: (parse_year(p.name) or 9999, p.name)):
        year = parse_year(folder.name)
        if not year or year > max_year:
            continue
        frames_dir = folder / "frames" / "cikis_fb"
        if not frames_dir.exists():
            continue
        frames = sorted(frames_dir.glob("*.png"))
        if len(frames) < 80:
            continue
        meta = load_clip_meta(folder)
        cases.append(
            {
                "folder": folder,
                "frames_dir": frames_dir,
                "title": meta.get("title") or re.sub(r"\s+\d{4}-.+$", "", folder.name),
                "trt_id": meta.get("trt_id") or "",
                "year": year,
                "tur": meta.get("tur") or "",
                "frame_total": len(frames),
            }
        )
        if len(cases) >= limit:
            break
    return cases


def sample_frames(tp: Any, frames_dir: Path, stride: int, start: int | None, end: int | None) -> list[Path]:
    frames = sorted(frames_dir.glob("*.png"), key=tp.frame_number)
    selected = []
    for idx, frame in enumerate(frames, start=1):
        n = tp.frame_number(frame)
        if start is not None and n < start:
            continue
        if end is not None and n > end:
            continue
        if ((idx - 1) % stride) == 0:
            selected.append(frame)
    return selected


def quiet_ocr_frames(tp: Any, pl: Any, frames: list[Path]) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for path in frames:
        img = pl.fp.rd(str(path))
        recs = pl.read_pos(img)
        lines = tp.clean_ocr_lines(recs)
        rows.append(
            {
                "frame": path.name,
                "frame_path": str(path),
                "frame_number": tp.frame_number(path),
                "lines": lines,
                "line_texts": [line["text"] for line in lines],
                "line_set": sorted(tp.line_set(lines)),
                "signature": tp.frame_signature(lines),
            }
        )
    return rows


def official_names(roles: dict[str, list[str]]) -> list[str]:
    out: list[str] = []
    for key in ("yonetmen", "yapimci", "senaryo_yazar", "oyuncular"):
        out.extend(roles.get(key) or [])
    return out


def is_single_token(tp: Any, name: str) -> bool:
    toks = [t for t in tp.fold(name).split() if t]
    return len(toks) < 2


def technical_leak(tp: Any, name: str) -> bool:
    f = tp.fold(name)
    markers = (
        "director",
        "producer",
        "screenplay",
        "music",
        "costume",
        "designed",
        "photography",
        "editor",
        "assistant",
        "attendant",
        "mayor",
        "master",
        "dresser",
        "artist",
        "camera",
        "composer",
        "written",
    )
    return any(m in f for m in markers)


def organization_or_address_leak(tp: Any, name: str) -> bool:
    f = tp.fold(name)
    markers = (
        "ltd",
        "limited",
        "inc",
        "company",
        "group",
        "studios",
        "studio",
        "productions",
        "laboratories",
        "corporation",
        "television",
        "road",
        "street",
        "avenue",
        "london",
        "england",
        "museum",
        "university",
        "railway",
        "city of",
        "all rights reserved",
        "dolby",
        "panavision",
        "metrocolor",
        "walt disney",
        "buena vista",
    )
    return any(m in f for m in markers)


def summarize_case(tp: Any, case: dict[str, Any], frame_rows: list[dict[str, Any]], events: list[dict[str, Any]], roles: dict[str, list[str]], elapsed: float) -> dict[str, Any]:
    names = official_names(roles)
    single = [n for n in names if is_single_token(tp, n)]
    tech = [n for n in names if technical_leak(tp, n)]
    org = [n for n in names if organization_or_address_leak(tp, n)]
    flags: list[str] = []
    if not any(row.get("lines") for row in frame_rows):
        flags.append("NO_TEXT")
    if not names:
        flags.append("NO_OFFICIAL_NAMES")
    if not roles.get("oyuncular"):
        flags.append("NO_CAST")
    if not roles.get("yonetmen"):
        flags.append("NO_DIRECTOR")
    if single:
        flags.append("SINGLE_TOKEN_OFFICIAL")
    if tech:
        flags.append("TECHNICAL_LEAK_OFFICIAL")
    if org:
        flags.append("ORG_OR_ADDRESS_LEAK_OFFICIAL")
    if len(roles.get("oyuncular") or []) > 25:
        flags.append("TOO_MANY_CAST")
    if sum(1 for e in events if e.get("type") == "transition_or_noise") > max(6, len(events) // 2):
        flags.append("MANY_TRANSITIONS")
    status = "OK" if not flags else ("FAIL" if "NO_TEXT" in flags or "SINGLE_TOKEN_OFFICIAL" in flags or "TECHNICAL_LEAK_OFFICIAL" in flags or "ORG_OR_ADDRESS_LEAK_OFFICIAL" in flags else "REVIEW")
    counts = {kind: sum(1 for e in events if e.get("type") == kind) for kind in ("static_card", "scroll_block", "transition_or_noise")}
    return {
        "status": status,
        "flags": flags,
        "folder": str(case["folder"]),
        "title": case["title"],
        "trt_id": case["trt_id"],
        "year": case["year"],
        "tur": case["tur"],
        "frame_total": case["frame_total"],
        "frames_ocrd": len(frame_rows),
        "text_frames": sum(1 for row in frame_rows if row.get("lines")),
        "events": len(events),
        "static_cards": counts["static_card"],
        "scroll_blocks": counts["scroll_block"],
        "transitions": counts["transition_or_noise"],
        "director_count": len(roles.get("yonetmen") or []),
        "producer_count": len(roles.get("yapimci") or []),
        "writer_count": len(roles.get("senaryo_yazar") or []),
        "cast_count": len(roles.get("oyuncular") or []),
        "ignored_count": len(roles.get("ignored") or []),
        "single_token_names": single,
        "technical_leak_names": tech,
        "organization_or_address_leak_names": org,
        "roles": roles,
        "elapsed_sec": round(elapsed, 2),
    }


def write_case_markdown(path: Path, summary: dict[str, Any], events: list[dict[str, Any]]) -> None:
    lines = [f"# {summary['title']} ({summary['year']})", ""]
    lines.append(f"- Status: {summary['status']}")
    lines.append(f"- Flags: {', '.join(summary['flags']) if summary['flags'] else '-'}")
    lines.append(f"- Frames OCR'd: {summary['frames_ocrd']} / {summary['frame_total']}")
    lines.append(f"- Events: {summary['events']} static={summary['static_cards']} scroll={summary['scroll_blocks']} transition={summary['transitions']}")
    lines.append("")
    lines.append("## Roles")
    for key, values in summary["roles"].items():
        lines.append(f"- {key}: {', '.join(values) if values else '-'}")
    lines.append("")
    lines.append("## Events")
    for event in events:
        lines.append(f"### {event['event_id']} {event['type']} {event['frame_start']}..{event['frame_end']}")
        lines.append(f"- Best frame: {event['best_frame']}")
        for text in event.get("lines") or []:
            lines.append(f"  - {text}")
        lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_reports(out_dir: Path, summaries: list[dict[str, Any]], payload: dict[str, Any]) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    (out_dir / "batch20_credit_structure_report.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    lines = ["# Batch20 Credit Structure Probe", ""]
    lines.append(f"- Films: {len(summaries)}")
    lines.append(f"- OK: {sum(1 for s in summaries if s['status'] == 'OK')}")
    lines.append(f"- REVIEW: {sum(1 for s in summaries if s['status'] == 'REVIEW')}")
    lines.append(f"- FAIL: {sum(1 for s in summaries if s['status'] == 'FAIL')}")
    lines.append("")
    lines.append("| status | year | title | roles | flags |")
    lines.append("|---|---:|---|---|---|")
    for s in summaries:
        roles = f"yon={s['director_count']} yap={s['producer_count']} sen={s['writer_count']} cast={s['cast_count']} ign={s['ignored_count']}"
        flags = ", ".join(s["flags"]) if s["flags"] else "-"
        lines.append(f"| {s['status']} | {s['year']} | {s['title']} | {roles} | {flags} |")
    lines.append("")
    lines.append("## Details")
    for s in summaries:
        lines.append(f"### {s['status']} - {s['title']} ({s['year']})")
        lines.append(f"- Flags: {', '.join(s['flags']) if s['flags'] else '-'}")
        lines.append(f"- Events: {s['events']} static={s['static_cards']} scroll={s['scroll_blocks']} transition={s['transitions']}")
        for key, values in s["roles"].items():
            lines.append(f"- {key}: {', '.join(values) if values else '-'}")
        lines.append("")
    (out_dir / "batch20_credit_structure_report.md").write_text("\n".join(lines), encoding="utf-8")

    try:
        from openpyxl import Workbook
    except Exception:
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    headers = [
        "status",
        "flags",
        "year",
        "title",
        "trt_id",
        "tur",
        "frames_ocrd",
        "text_frames",
        "events",
        "static_cards",
        "scroll_blocks",
        "transitions",
        "director_count",
        "producer_count",
        "writer_count",
        "cast_count",
        "ignored_count",
        "yonetmen",
        "yapimci",
        "senaryo_yazar",
        "oyuncular",
        "ignored",
        "single_token_names",
        "technical_leak_names",
        "organization_or_address_leak_names",
        "elapsed_sec",
        "folder",
    ]
    ws.append(headers)
    for s in summaries:
        row = []
        for h in headers:
            if h in ("yonetmen", "yapimci", "senaryo_yazar", "oyuncular", "ignored"):
                value = s["roles"].get(h) or []
            else:
                value = s.get(h)
            row.append(json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value)
        ws.append(row)

    ev = wb.create_sheet("Events")
    ev.append(["title", "year", "event_id", "type", "frame_start", "frame_end", "best_frame", "lines"])
    for film in payload["films"]:
        for event in film["events_detail"]:
            ev.append([
                film["summary"]["title"],
                film["summary"]["year"],
                event.get("event_id"),
                event.get("type"),
                event.get("frame_start"),
                event.get("frame_end"),
                event.get("best_frame"),
                json.dumps(event.get("lines") or [], ensure_ascii=False),
            ])
    wb.save(out_dir / "batch20_credit_structure_report.xlsx")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=20)
    parser.add_argument("--max-year", type=int, default=2010)
    parser.add_argument("--stride", type=int, default=3)
    parser.add_argument("--start", type=int, default=None)
    parser.add_argument("--end", type=int, default=None)
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    args.out_dir.mkdir(parents=True, exist_ok=True)
    tp = load_module("temporal_probe_batch20", TEMPORAL_PROBE)
    pl = tp._load_pipeline100()
    cases = discover_cases(args.limit, args.max_year)
    if len(cases) < args.limit:
        print(f"[warn] only {len(cases)} cases found", flush=True)

    films: list[dict[str, Any]] = []
    summaries: list[dict[str, Any]] = []
    started = time.perf_counter()
    for idx, case in enumerate(cases, start=1):
        t0 = time.perf_counter()
        frames = sample_frames(tp, case["frames_dir"], args.stride, args.start, args.end)
        print(f"[{idx}/{len(cases)}] {case['title']} ({case['year']}) frames={len(frames)}", flush=True)
        frame_rows = quiet_ocr_frames(tp, pl, frames)
        events = tp.segment_frames(frame_rows)
        roles = tp.parse_roles(events)
        elapsed = time.perf_counter() - t0
        summary = summarize_case(tp, case, frame_rows, events, roles, elapsed)
        safe = re.sub(r"[^A-Za-z0-9._-]+", "_", f"{idx:02d}_{case['year']}_{case['title']}").strip("_")
        film_dir = args.out_dir / "films" / safe
        film_dir.mkdir(parents=True, exist_ok=True)
        (film_dir / "result.json").write_text(
            json.dumps({"summary": summary, "events": events, "frame_ocr": frame_rows}, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        write_case_markdown(film_dir / "result.md", summary, events)
        summaries.append(summary)
        films.append({"summary": summary, "events_detail": events, "result_dir": str(film_dir)})
        print(f"    -> {summary['status']} flags={','.join(summary['flags']) or '-'} roles cast={summary['cast_count']} dir={summary['director_count']} elapsed={elapsed:.1f}s", flush=True)

    payload = {
        "config": {
            "limit": args.limit,
            "max_year": args.max_year,
            "stride": args.stride,
            "start": args.start,
            "end": args.end,
            "elapsed_sec": round(time.perf_counter() - started, 2),
        },
        "summary_counts": {
            "films": len(summaries),
            "ok": sum(1 for s in summaries if s["status"] == "OK"),
            "review": sum(1 for s in summaries if s["status"] == "REVIEW"),
            "fail": sum(1 for s in summaries if s["status"] == "FAIL"),
        },
        "films": films,
    }
    write_reports(args.out_dir, summaries, payload)
    print(args.out_dir / "batch20_credit_structure_report.md")
    print(args.out_dir / "batch20_credit_structure_report.xlsx")
    print(args.out_dir / "batch20_credit_structure_report.json")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
