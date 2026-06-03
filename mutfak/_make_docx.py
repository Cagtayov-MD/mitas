# -*- coding: utf-8 -*-
"""MITAS Model Envanteri -> Word (.docx). Kompakt, belirgin model basliklari."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.table import WD_TABLE_ALIGNMENT

NAVY  = RGBColor(0x16,0x31,0x5C)
WHITE = RGBColor(0xFF,0xFF,0xFF)
LABEL = RGBColor(0x33,0x44,0x5E)
EXCLR = RGBColor(0x4A,0x55,0x68)
BODY  = RGBColor(0x22,0x2B,0x36)
GREY  = RGBColor(0x6B,0x7A,0x92)
FONT  = "Segoe UI"

def shade(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),fill)
    pPr.append(shd)

def bottom_border(par, color, sz=12):
    pPr = par._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'),'single'); b.set(qn('w:sz'),str(sz)); b.set(qn('w:space'),'2'); b.set(qn('w:color'),color)
    pbdr.append(b); pPr.append(pbdr)

doc = Document()
s = doc.sections[0]
s.page_width = Cm(21.0); s.page_height = Cm(29.7)
s.top_margin = Cm(1.3); s.bottom_margin = Cm(1.3); s.left_margin = Cm(1.5); s.right_margin = Cm(1.5)
nrm = doc.styles['Normal']; nrm.font.name = FONT; nrm.font.size = Pt(9.5); nrm.font.color.rgb = BODY
pf = nrm.paragraph_format; pf.space_before = Pt(0); pf.space_after = Pt(1); pf.line_spacing = 1.0

def title_block():
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("MITAS — Model Envanteri"); r.font.bold=True; r.font.size=Pt(20); r.font.color.rgb=NAVY; r.font.name=FONT
    p2 = doc.add_paragraph(); p2.paragraph_format.space_after = Pt(6)
    r2 = p2.add_run("Sistemde kullanılan ve planlanan yapay zeka modelleri · 3 Haziran 2026"); r2.font.size=Pt(9); r2.font.color.rgb=GREY; r2.font.italic=True
    p3 = doc.add_paragraph(); p3.paragraph_format.space_after = Pt(8)
    r3 = p3.add_run("Her model için ne işe yaradığı, nasıl çalıştığı, bir örnek ve çözdüğü sorun verilmiştir. ASR, OCR ve Çeviri modülleri çalışır durumdadır; Müzik, Yüz ve Sahne tanıma planlı olduğu için sonda yer alır.")
    r3.font.size=Pt(9); r3.font.color.rgb=BODY

def sec_title(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(9); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
    r = p.add_run(t); r.font.bold=True; r.font.size=Pt(13); r.font.color.rgb=NAVY; r.font.name=FONT
    bottom_border(p,'1F4A7A',10)

def sec_desc(t):
    p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    r = p.add_run(t); r.font.italic=True; r.font.size=Pt(8.6); r.font.color.rgb=GREY

def model(name):
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2); p.paragraph_format.keep_with_next=True
    shade(p,'16315C')
    r = p.add_run("  " + name + "  "); r.font.bold=True; r.font.size=Pt(11); r.font.color.rgb=WHITE; r.font.name=FONT

def field(label, text):
    p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.left_indent=Cm(0.2); p.paragraph_format.keep_with_next=True
    rl = p.add_run(label + " "); rl.font.bold=True; rl.font.size=Pt(9); rl.font.color.rgb=LABEL; rl.font.name=FONT
    rt = p.add_run(text); rt.font.size=Pt(9.3); rt.font.color.rgb=BODY; rt.font.name=FONT

def ex(text):
    pl = doc.add_paragraph(); pl.paragraph_format.space_after=Pt(0); pl.paragraph_format.left_indent=Cm(0.2); pl.paragraph_format.keep_with_next=True
    rl = pl.add_run("Örnek:"); rl.font.bold=True; rl.font.size=Pt(9); rl.font.color.rgb=LABEL; rl.font.name=FONT
    p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(1); p.paragraph_format.left_indent=Cm(0.5)
    for i, ln in enumerate(text.split("\n")):
        r = p.add_run(ln); r.font.name="Consolas"; r.font.size=Pt(8.3); r.font.color.rgb=EXCLR
        if i < len(text.split("\n")) - 1:
            r.add_break()

DATA = [
 ("1 · ASR — Konuşmayı Yazıya Dönüştürme",
  "Sesi alır, kim ne zaman ne söyledi sorusuna zaman damgalı metinle cevap verir. Sistemin en olgun modülü.",
  [
   ("faster-whisper · large-v3-turbo",
    "Ana konuşma tanıma motoru; varsayılan hızlı yolda sesi metne çevirir.",
    "OpenAI Whisper mimarisinin hızlandırılmış (int8/float16) hâli. Encoder-decoder bir yapay sinir ağı; 16 kHz sesi dinler ve dikkat mekanizmasıyla her ses parçasını metne, dil etiketine ve güven skoruna çevirir. Turbo sürümü tam modelin küçültülmüş hâlidir — gerçek zamanın yaklaşık 25 katı hızda çalışır.",
    "16 kHz mono ses → (0.5–3.2sn) \"İyi akşamlar, gündemdeki gelişmeler...\" | dil: tr | güven: 0.94",
    "Türkçe ve çok dilli arşivi hızlı, makul maliyetle yazıya dökmek."),
   ("faster-whisper · large-v3",
    "Kalite öncelikli ağır model; hızlı model takılırsa veya şüpheli sonuç verirse devreye girer.",
    "Turbo ile aynı mimari ama tam boy decoder (yaklaşık 1.5 milyar parametre); daha geniş bağlam, daha düşük hata. Tüm dosyayı baştan çözmez — sadece riskli işaretlenen parçaları yeniden çözüp temiz transcript'in içine yamalar (seçici onarım).",
    "Tetik: bir cümlede erken bitiş şüphesi → sadece o 8 saniyelik parça large-v3 ile yeniden çözülür → eksik cümle geri gelir",
    "Hızı korurken kritik anlarda kelime kaybını ve hatayı engellemek."),
   ("Silero VAD",
    "Seste konuşma var mı, varsa nerede — onu bulur (Voice Activity Detection).",
    "Hafif bir sinir ağı, sesi 30 milisaniyelik küçük pencerelere bölüp her birine \"burada konuşma var mı?\" olasılığı verir. Eşiği geçen ardışık pencereleri birleştirip konuşma aralıkları listesi çıkarır.",
    "60 dakikalık ses → konuşma aralıkları [ (0.5–3.2sn), (5.1–8.8sn), ... ] (sessiz/müzik bölgeleri ayıklanır)",
    "Sessiz ve müzikli bölgeleri ASR'a sokmayarak hayalet metni ve boşa zaman harcamayı önlemek."),
   ("WhisperX",
    "Cümle bazlı zaman damgasını kelime bazına indirir — her kelimenin tam kaçıncı saniyede söylendiğini bulur.",
    "Whisper'ın ürettiği cümle metnini ses dalgasıyla hizalar (forced-alignment). Her kelimeye başlangıç/bitiş ve güven skoru atar. Bağımlılık çakışmasını önlemek için ayrı bir ortamda alt-süreç olarak çağrılır.",
    "\"merhaba dünya\" (0.5–1.8sn) → merhaba (0.50–0.92) · dünya (0.95–1.80)",
    "Altyazı senkronu ve \"bu kelimeye tıkla, videoda oraya git\" türü kelime düzeyinde arama."),
   ("pyannote-audio · konuşmacı ayrımı",
    "Konuşmacı ayrımı yapar: kim, ne zaman konuştu (diarization).",
    "Her konuşmacının ses imzasını çıkarır ve benzer imzaları kümeleyerek SPEAKER_00, SPEAKER_01 gibi etiketler verir. Güven eşiğinin altında kalan parçaya etiket basmaz — yanlış atama yapmaktansa boş bırakmak ilkesi.",
    "2 sunuculu haber bülteni → SPEAKER_00 (0.2–4.5sn), SPEAKER_01 (5.0–9.3sn) → her satıra konuşmacı eklenir",
    "Panel, röportaj, haber gibi çok kişili içerikte \"kim söyledi\" bilgisini üretmek."),
   ("SpeechBrain · dil tanıma (VoxLingua107)",
    "Sesin hangi dilde olduğunu tespit eder (107 dil: tr, ar, az, ku, en...).",
    "Sesten dil imzası çıkaran bir sinir ağı; kısa pencerelere bakıp 107 dil üzerinde olasılık dağılımı verir. Klibin ana dilini, pencerelerin yarısından fazlasına hakim olan dil olarak belirler. Şu an sadece ölçüm/raporlama yapıyor; otomatik yönlendirme henüz kapalı.",
    "20 pencerelik ses → { tr: 0.78, ar: 0.12, en: 0.10 } → ana dil: Türkçe",
    "Arşivdeki Arapça/Azerice/Kürtçe içeriği otomatik ayırmak ve doğru çeviri rotasına hazırlamak."),
   ("DeepFilterNet · ses temizleme",
    "Arka plan gürültüsünü/müziğini konuşmadan ayırarak ASR doğruluğunu artırır.",
    "Derin öğrenme tabanlı gürültü bastırma filtresi; zaman-frekans maskesiyle gürültü bileşenini söker. Şu an profil ayarlarında işaretli (stüdyo/film/belgesel/spor) ama pipeline'a henüz bağlanmadı — niyet kayıtlı, çağrı bekliyor.",
    "Arka planda müzik olan film diyaloğu → sadece konuşma kalan temiz ses",
    "Stüdyo dışı, gürültülü arşiv kaydında ASR hatalarını azaltmak."),
   ("Özel İsim Düzeltme",
    "ASR'ın ses benzerliği yüzünden yanlış yazdığı bilinen özel adları doğru biçimine düzeltir.",
    "Yapay zeka değil, \"yanlış → doğru\" kural listesi. Her segment metninde arar ve düzeltir. Orijinal metni silmez, düzeltilmiş hâli ayrı bir alanda durur (incelenebilir ikinci katman).",
    "\"Barış Mango konseri\" → \"Barış Manço konseri\"",
    "Tahmin edilebilir özel ad hatalarını arşiv aramasını bozmadan, güvenli biçimde gidermek."),
   ("Tahmini Kelime Zamanlama",
    "WhisperX başarısız olursa kelime zaman damgalarını yine de üretir (yedek mekanizma).",
    "Model değil basit hesap: cümlenin süresini kelime sayısına eşit aralıklara böler ve her kelimeye yaklaşık bir başlangıç/bitiş verir.",
    "\"merhaba dünya\" (0.5–1.8sn) → merhaba (0.50–1.15) · dünya (1.15–1.80)",
    "Hizalama motoru çökse bile her çıktının kelime zamanına sahip olmasını garanti etmek."),
  ]),
 ("2 · Çeviri — Yabancı Dili Türkçeye",
  "ASR'dan çıkan yabancı dil satırlarını Türkçeye çevirir. Çalışıyor ve sisteme entegre.",
  [
   ("OPUS-MT · İngilizce → Türkçe",
    "İngilizceden Türkçeye hızlı ve kaliteli çeviri uzmanı.",
    "Encoder-decoder çeviri modeli; metni alt-kelime parçalarına bölüp hedef dile çevirir. Yalnız İngilizce → Türkçe yapar; diğer diller çok dilli modele yönlendirilir. Düşük bellekte hızlı çalışır.",
    "\"The parliament voted on the budget.\" → \"Meclis bütçe üzerinde oy kullandı.\"",
    "En sık karşılaşılan İngilizce kaynakları hızlı ve ucuz çevirmek."),
   ("NLLB-200-3.3B · çok dilli çeviri",
    "İngilizce dışı tüm dillerden (özellikle Arapça) Türkçeye çeviren çok dilli ana model.",
    "200 dili destekleyen tek çeviri modeli. Kaynak ve hedef dil kodlarıyla belirtilir (örn. Arapça → Türkçe) ve hedef dil zorlanarak çeviri üretilir.",
    "\"البرلمان صوت على الميزانية\" → \"Meclis bütçeyi oyladı.\"",
    "Tek modelle 200+ dilden Türkçeye çeviri — Arapça başta olmak üzere bölgesel diller."),
   ("Arapça Lehçe Normalizasyonu",
    "Suriye/Mısır/Körfez/Mağrip lehçelerini standart Arapçaya çevirerek çeviri kalitesini yükseltir.",
    "Lehçeye özgü sinyal kelimeleri sayarak lehçeyi tespit eder, sonra kural tabanlı değiştirmeyle standart Arapçaya çevirir. Bu temiz metin çeviri modeline gönderilir.",
    "\"هلق وين رايح؟\" (Levanten) → \"الآن أين تذهب؟\" (standart) → ardından Türkçeye çevrilir",
    "Çeviri modelinin standart Arapça ağırlıklı eğitiminden kaynaklanan lehçe zayıflığını gidermek."),
   ("NLLB-200-distilled-1.3B (yedek)",
    "3.3B modelin küçültülmüş, hızlı versiyonu; kurulu tutuluyor ama şu an rotaya bağlı değil.",
    "Aynı mimari, daha küçük parametre. Kalitesi 3.3B'nin biraz altında, hızı belirgin yüksek.",
    "Aynı Arapça girdi → daha hızlı çeviri, hafif kalite farkı",
    "3.3B çok yavaş kalırsa veya bellek yetmezse devreye alınacak hız yedeği."),
  ]),
 ("3 · OCR & Görsel-Dil — Ekrandaki Yazıyı Okuma",
  "Jenerik, künye, alt yazı ve tabelaları okur. Akan/duran jeneriği tek görüntüye dönüştüren yöntemler ayrı belgededir.",
  [
   ("PaddleOCR (birincil)",
    "Ana OCR motoru; karelerden metin kutularını bulup içlerini okur.",
    "İki aşamalı: önce bir tespit modeli görüntüdeki yazı bölgelerini kutu olarak işaretler, sonra bir tanıma modeli her kutunun içindeki yazıyı okur. GPU üzerinde çalışır; çok uzun jenerik görüntülerini dikey dilimlere bölerek işler.",
    "Jenerik karesi → [ \"YÖNETMEN\" kutu(x,y,w,h) güven 0.94 ], [ \"AHMET TEKİN\" ... ]",
    "Büyük ölçekli, çok satırlı jeneriklerde hızlı ve doğru metin tespiti ve okuma."),
   ("OneOCR (yedek)",
    "Windows'un kendi OCR motoru; birincil motor bozuk metin ürettiğinde devreye giren akıllı yedek.",
    "Windows yerleşik OCR altyapısını çağırır; görüntüyü alır, satır satır metin + kutu + güven döner. Pipeline'a yalnız gerektiğinde çağrılır.",
    "Birleştirilmiş jenerik görüntüsü → temiz satırlar [ \"YÖNETMEN\", \"GÖRÜNTÜ YÖNETMENİ\", ... ]",
    "Birincil motorun çuvalladığı zor görüntülerde ikinci bir okuma şansı vermek."),
   ("qwen2.5vl — \"Bekçi / GÖZ\"",
    "OCR yapmaz, karar verir: \"Bu kare jenerik mi, film görüntüsü mü?\" Yanlış görüntünün OCR'a gitmesini engeller.",
    "Tek kareye bakan bir görsel-dil modeli; yapılandırılmış cevap döner (jenerik mi, akıyor mu, arka plan hareketli mi, kaç sütun, hangi yöntem önerilir). Hareket ölçümünü matematik yapar, görünüm yorumunu model yapar, kod ikisini birleştirip doğru motoru seçer.",
    "akan jenerik karesi → { jenerik: evet, hareket: kayan, arka plan: hareketli sahne }\nkadın yüzü karesi → { jenerik: hayır } → bu kare OCR'a gönderilmez",
    "Projenin en kritik açığı: film görüntüsünün \"jenerik\" sanılıp çöp üretmesi. Testte 5/5 doğru ayırdı."),
   ("GLM-OCR",
    "Görsel-dil tabanlı OCR; özellikle kırpılmış küçük metin kutularını hızlı ve temiz okur. Karşılaştırmanın kazananı.",
    "Yerel olarak çalışır; görüntü + komut alır, yazıyı döndürür. Düşük sıcaklık ve tekrar cezasıyla döngüye girmeden kararlı okur. Tam sayfa veya küçük kırpıntı kabul eder.",
    "Bir metin kutusunun 3× büyütülmüş kırpıntısı → \"YÖNETMEN\" (yaklaşık 0.3–1.0 sn)",
    "Hızlı, kurulumu kolay ve OneOCR ile çapraz doğrulamaya uygun bir OCR ikinci görüşü."),
   ("InternVL3-8B",
    "Zor Türkçe jenerikleri en iyi okuyan görsel-dil modeli (kıyaslamada birinci).",
    "Görüntüyü dinamik karolara bölüp işleyen bir model. \"Ekrandaki her şeyi satır satır yaz\" komutuyla tam sayfa metni döndürür.",
    "Bulanık Türkçe jenerik sayfası → \"YÖNETMEN / AHMET TEKİN / SENARYO / ...\" (Ş,Ğ,İ doğru)",
    "Diğer modellerin kaçırdığı Türkçe karakterleri (Ş, Ğ, İ) zor koşullarda doğru okumak."),
   ("LocateAnything-3B",
    "OCR yapmaz; görüntüde metnin nerede olduğunu (kutu koordinatları) bulur. GLM-OCR ile zincir kurulur.",
    "\"Tüm metni kutu olarak tespit et\" komutuyla koordinatlar döndüren bir model. Bulduğu her kutu kırpılıp ayrıca okutulur.",
    "Jenerik karesi → 8–12 kutu koordinatı → her kutu kırpılıp okunur → \"YÖNETMEN\": 0.95",
    "Tam sayfa yerine sadece yazı bölgelerini yüksek çözünürlükte okuyarak doğruluğu artırmak."),
   ("MiniCPM-V 2.6",
    "Tek görüntüden Türkçe jenerik metni çıkaran görsel-dil modeli (kıyaslamada üçüncü).",
    "Çok modlu model; görüntü ve komutu birlikte alıp metin döndürür. Mevcut donanımda çalışır.",
    "Jenerik sayfası → \"YÖNETMEN BARIŞ PİRHASAN / SENARYO ...\"",
    "Genel amaçlı VLM-OCR alternatifi; kazananların yanında karşılaştırma noktası."),
   ("Tesseract (planlı)",
    "Klasik, kütüphanesiz OCR alternatifi; karşılaştırma adayı.",
    "Karakter segmentasyonu + dil modeliyle çalışan köklü açık kaynak OCR. Kelime düzeyinde kutu ve güven üretir. Motor listesinde tanımlı ama henüz üretimde aktif değil.",
    "Kart görüntüsü → [ \"KAMERA\" kutu güven 0.87 ]",
    "Platform bağımsız, kütüphane gerektirmeyen yedek OCR ve karşılaştırma referansı."),
  ]),
 ("4 · Üst Denetim & Özet — Büyük Dil Modelleri",
  "ASR/OCR çıktısının üstünde çalışan akıl katmanı: hata/şüphe yakalar, kanıt ister, özet çıkarır. \"Hâkim değil, kanıt isteyen denetçi\" ilkesiyle tasarlandı.",
  [
   ("Claude (Sonnet) · özet motoru",
    "Transcript'ten özet ve film/dizi künye metni üretir.",
    "Büyük dil modeli; içerik profiline göre (spor / genel) farklı komut setiyle çağrılır. Yerel modeller özet formatını tutmadığı için bu işin birincil çözümü buraya bağlandı; yerel model ve özetleyici yedek zinciri arkada durur.",
    "40 dk transcript → \"Önemli Konular: 1)... 2)... · Geçen İsimler: ...\"",
    "Tutarlı kalitede Türkçe özet — yerel modellerin başaramadığı görev."),
   ("Kimi-K2.6 (güçlü aday)",
    "Üst denetimde en güçlü yeni aday; özel isim/bağlam testinde diğerlerini geçti.",
    "Büyük dil modeli. Transcript'i değiştirmez; \"şu isim şüpheli\", \"bu bir tabela, dokunma\", \"tarihsel olarak imkânsız\" gibi yapılandırılmış bulgular üretir. Zor vakaları doğru sınıflandırdı.",
    "\"C. AUBREY SMlTH\" (OCR hatası) → { aksiyon: sadece-bağlantı-kur, kanonik: \"C. Aubrey Smith\" }",
    "Yerel isim/tabela tuzaklarına düşmeden güvenilir denetim kararı."),
   ("Gemma-4-26b (varsayılan aday)",
    "Üst denetimin mevcut varsayılan metin modeli (hız/kalite dengesi iyi).",
    "Büyük dil modeli. İzinli aksiyonlar listesi içinde yanıt verir (bağlantı kur, düzeltme öner, kırmızı bayrak, kare iste...) ve doğrudan metni değiştirmez; sisteme öneri döndürür.",
    "\"1936 filminde Ahmet Necdet Sezer oyuncuydu\" → { tarihsel imkânsızlık, kaynak doğrulaması iste }",
    "100.000 video ölçeğinde insan onayı olmadan güvenli, izlenebilir denetim."),
   ("Qwen3.6-27b (kalite / ikinci görüş)",
    "En dengeli bağlam ayrımı; belirsiz vakalarda ikinci görüş/hakem.",
    "Yoğun büyük model. \"Normal konuşma mı, alıntı mı, tabela mı?\" ayrımını en tutarlı yapan model; ancak yavaş olabiliyor.",
    "\"Barış Mango geçiyor\" → normal konuşmaysa düzeltme öner, ekranda tabelaysa sadece incele",
    "Diğer modelin kaçırdığı ince bağlam kırıklarını yakalamak."),
   ("Qwen3-30b-a3b (hızlı alternatif)",
    "Qwen 27B'nin daha hızlı alternatifi.",
    "Seyrek-uzman model; yoğun modele yakın kalite, belirgin hız avantajı. Bazı sınır vakalarında kaçırma yapabiliyor.",
    "Uzun belgesel bloğu → Qwen 27B'ye yakın kalite, yaklaşık %30 daha hızlı",
    "Üretimde verim/kalite dengesini tutmak."),
   ("Llama-3.3-70b (denetçi / hakem)",
    "Yüksek riskli kararların son hakemi; çok temkinli çalışır.",
    "Büyük model. Doğru komutla \"düzeltme önermekten kaçınan\", emin değilse incelemeye yönlendiren konservatif bir hakem. Maliyeti yüksek olduğu için ana akışta değil, doğrulama turunda.",
    "Başka model \"değiştir\" dediğinde Llama \"güven düşük → sadece incele\" diyerek frene basabilir",
    "Otomatik değiştirmenin ikinci güvenlik kilidi."),
   ("Nemotron-3-nano-omni (planlı · görsel kanıt)",
    "Üst denetim kare/kırpıntı kanıtı istediğinde görüntüyü okuyup metin üreten çok modlu model.",
    "Görüntü + metni birlikte işler. Denetim ikinci turunda \"şu kareyi/künyeyi oku\" aksiyonundan gelen görseli okuyarak kanıt sağlar.",
    "\"künyeyi tekrar oku\" aksiyonu → künye kırpıntısı → \"BESTE: Selâhattin Pınar\"",
    "Metin denetiminin görsel kanıta ihtiyaç duyduğu vakalar."),
   ("qwen3-vl · video soru-cevap",
    "Videoya soru sorup cevap almak için kullanılan görsel-dil modeli (deneysel araç).",
    "Çıkarılan kareler + soru modele verilir. Düşünme modu kapalı çalıştırılır (aksi halde döngüye giriyor), tekrar cezasıyla stabilize edilir.",
    "Kare + \"Bu sahnede kim/ne var?\" → sahne/karakter tanımı",
    "Görsel doğrulama ve sahne soru-cevap için yardımcı araç."),
  ]),
 ("5 · Müzik Tanıma — Planlı (v0.2 hedefi)",
  "Müzik programlarında \"hangi eser, kim seslendirdi\" sorusunu çözecek katman. Hedef: zamanla büyüyen TRT Müzik Kütüphanesi.",
  [
   ("Repertuvar Bankası (çekirdek)",
    "Eser ve sanatçı varlık sözlüğü — Müzik Kütüphanesi'nin temeli. Okunan adları doğru kanonik kayda bağlar.",
    "İki tablo: Eser (kanonik ad, takma adlar, besteci, makam, söz) ve Sanatçı (ad, roller). Türkçe karaktere duyarlı bulanık eşleştirmeyle sorgulanır. Seed: TRT Repertuvar Kurulu verisi + MusicBrainz Türkçe alt kümesi.",
    "sorgu(\"ne güzel güldün\") → { eser: \"Ne Güzel Güldün\", besteci: \"Selâhattin Pınar\" }",
    "Parmak izi eşleşmese bile künye/anonstan eser kimliği üretmek; kütüphanenin büyümesini sağlamak."),
   ("ASR Anons Parser",
    "Sunucunun \"şimdi sizlere X'ten Y'yi dinleteceğiz\" anonsunu yakalayıp eser+sanatçı çıkarır.",
    "Müzik bölümünden önceki yaklaşık 30 saniyelik transcript'te Türkçe anons kalıplarını arar; eşleşen kısmı Repertuvar Bankası ile eşleştirip aday üretir.",
    "\"Şimdi Bülent Ersoy'dan Ne Güzel Güldün'ü dinleteceğiz\" → { sanatçı: Bülent Ersoy, eser: Ne Güzel Güldün }",
    "Künye yokken veya künyeden önce en erken müzik kimlik sinyalini yakalamak."),
   ("PANNs CNN14 · ses sınıflandırma",
    "Sesi sınıflandırır: konuşma / müzik / alkış / sessizlik. Müziğin nerede başlayıp bittiğini bulur.",
    "Geniş ses veri kümesinde eğitilmiş bir sinir ağı. 1 saniyelik pencerelerle sınıflandırır, sonra başlangıç/bitiş sınırlarını tespit eder. Şarkı ve alkış alt sınıfları müzik programı için kritik.",
    "90 sn klip → [ konuşma 0–12sn, müzik 12–74sn, alkış 74–79sn ]",
    "Müzik/konuşma sınırını yaklaşık ±2 sn toleransla bulup performans segmentini işaretlemek."),
   ("Demucs · vokal ayırma (ileri aşama)",
    "Müzikte vokali enstrümandan ayırır; sonra vokal üzerinde ASR koşup şarkı sözünden eseri bulur.",
    "Kaynak ayırma modeli sesi vokal/davul/bas/diğer olarak böler. Ayrılan vokal ASR'a verilir, çıkan söz şarkı sözü veritabanıyla eşleştirilir.",
    "90 sn müzik → vokal.wav → ASR → \"Ağlasam mı gülsem mi\" → eser eşleşmesi",
    "Künye ve anons yokken şarkı sözünden eser tespiti."),
   ("Chromaprint + AcoustID · parmak izi (doğrulayıcı)",
    "Akustik parmak izi: bir kayıt daha önce tanımlandıysa eşleştirir. Birincil değil, doğrulayıcı katman.",
    "Sesten kompakt bir parmak izi üretir; bu izi yerel veritabanına veya dış servise (yalnız iz gider, ses gitmez) sorar. Canlı stüdyo icrasında eşleşme beklenmez; arşiv kayıt tekrarında işe yarar. Dış servis için kurum onayı bekleniyor.",
    "Arşiv kaydı çalıyor → iz → yerel veritabanı → { eşleşme: \"Barış Manço - Gülpembe\", 0.97 }",
    "Künye/anons varsa onları doğrulayan ek güven sinyali."),
  ]),
 ("6 · Altyapı Araçları",
  "Tüm modülleri besleyen, doğrudan \"yapay zeka\" olmayan ama olmazsa hiçbir şeyin çalışmadığı temel araçlar.",
  [
   ("ffmpeg",
    "Medya İsviçre çakısı: videodan ses çıkarır, kare çıkarır, format dönüştürür, kanal ayırır. Her modülün giriş kapısı.",
    "Komut satırı aracı; çağrılır, sonucu geçici dosyaya yazar, modüller o dosyayı okur. Her codec'i (arşivin eski formatları dahil) açabilir.",
    "Ses çıkarma: film.mp4 → 16 kHz mono ses (ASR'ın beklediği format)\nKare çıkarma: film.mp4 → saniyede 2.5 kare, 640×360 görüntü (OCR/yüz/sahne için)",
    "Doğrudan okunamayan arşiv formatlarını tüm modellerin anlayacağı standart girdiye çevirmek."),
   ("ffprobe",
    "Medya dosyasının künyesini okur: süre, codec, kanal sayısı, kare hızı.",
    "Dosyayı açar, başlık bilgisini döndürür (içeriği çözmeden). Pipeline hangi profili/modeli seçeceğine bunun verisine bakarak karar verir; süre doğrulaması zorunlu adımdır.",
    "ffprobe film.mp4 → { süre: 5432.17sn, kanal: 2, codec: h264, örnekleme: 48000 }",
    "İşlemeden önce medyanın ne olduğunu bilmek."),
   ("PyTorch + CUDA",
    "GPU üzerinde derin öğrenme hesaplamasının temel altyapısı.",
    "Modeller GPU belleğine yüklenip tensor hesapları kartta yapılır. MITAS modelleri sıralı yükler (biri biter, bellekten boşaltılır, diğeri yüklenir) — 24 GB sınırını aşmamak için.",
    "Mevcut GPU'da faster-whisper → gerçek zamanın yaklaşık 25 katı hız",
    "CPU'da dakikalar sürecek işi saniyelere indirmek; 1 saatlik klip hedefini mümkün kılmak."),
   ("CTranslate2",
    "ASR ve çeviri modellerini düşük bellekte, hızlı çalıştıran çıkarım motoru.",
    "Modelleri sıkıştırılmış (int8/float16) olarak çalıştırır. faster-whisper ve çeviri modelleri bu motorun üzerinde koşar.",
    "large-v3-turbo'yu yaklaşık 1–2 GB bellekle çalıştırır (resmi sürüme göre yaklaşık 4× hız)",
    "GPU belleğini verimli kullanıp hızı katlamak."),
   ("ONNX Runtime",
    "ONNX formatındaki modelleri GPU'da çalıştırır (özellikle yüz tanıma modellerinin motoru olacak).",
    "Model grafiğini optimize edip GPU üzerinde çalıştırır; daha küçük bellek ve hızlı başlangıç sağlar.",
    "Yüz tanıma modeli → GPU'da toplu kimlik vektörü çıkarımı",
    "ONNX modelleri için hafif, hızlı çalıştırma katmanı."),
   ("OpenCV",
    "Tüm görüntü işleme: kare okuma/yazma, hız ölçümü, hareket takibi, keskinlik.",
    "Görüntü kütüphanesi. Faz korelasyonu (kayma hızı), optik akış (hareket), keskinlik ölçümü, yazı maskesi gibi işlemleri sağlar. Türkçe-İ içeren Windows yollarında özel okuma yöntemiyle çalışır.",
    "İki ardışık kare → \"yazı 8.4 piksel/kare kayıyor\"",
    "GPU gerektirmeyen, OCR yöntemlerinin dayandığı görüntü işleme tabanı."),
   ("Modüler Sanal Ortam İzolasyonu",
    "Her modül kendi izole ortamında çalışır; biri diğerini bozamaz.",
    "ASR, OCR, ses sınıflandırma gibi çakışan bağımlılıklar ayrı ortamlarda yaşar. Modüller birbiriyle doğrudan değil, alt-süreç veya servis üzerinden konuşur.",
    "WhisperX ayrı ortamda → ASR onu alt-süreç olarak çağırıp sonucu alır",
    "Bir kurulumun diğer modülleri bozmasını önlemek; her parçayı bağımsız güncellenebilir kılmak."),
  ]),
 ("7 · Yüz Tanıma — Planlı (v0.3 hedefi)",
  "Henüz yazılmadı; kütüphaneler kurulu, plan hazır. Filmde kimlik kapalı (sadece \"yüz var\"), haber/program/belgeselde önemli kişi takibi.",
  [
   ("InsightFace / buffalo_l",
    "Tek pakette hem yüz tespiti hem yüz tanıma (kimlik vektörü) sağlayan çerçeve.",
    "Hazır modeller (tespit + tanıma) tek paket olarak gelir. Tek çağrıyla hem yüz konumunu hem 512 boyutlu kimlik imzasını döndürür. Lisans onayı kurum hukukunda bekliyor.",
    "Kare → [ { kutu, güven 0.94, kimlik vektörü (512 sayı) } ]",
    "İki ayrı model yükleme ve senkron yükünü ortadan kaldırmak."),
   ("SCRFD · yüz tespiti",
    "Karelerde yüzleri hızla tespit eder ve 5 nirengi noktası verir.",
    "Hızlı bir yüz dedektörü; büyük ve küçük yüzleri ayrı katmanlarda örnekler. İki geçişte kullanılır: önce hızlı tarama (yüz var mı?), sonra yoğun tespit.",
    "1080p kare → [ { kutu [420,180,612,410], güven 0.92, nirengi noktaları } ]",
    "1 saatlik klibi sadece yüz olan anlara odaklanarak hızlı taramak."),
   ("ArcFace · kimlik vektörü",
    "Hizalanmış yüzü 512 boyutlu kimlik vektörüne çevirir; kişiler bu vektörle karşılaştırılır.",
    "Eğitimde aynı kişiyi yakın, farklı kişiyi uzak vektörlere yerleştiren bir model. Karşılaştırma benzerlik skoruyla yapılır (eşik üstü ise aynı kişi).",
    "112×112 hizalı yüz → 512 sayılık vektör → bankadaki vektörlerle kıyas",
    "Işık, poz ve yaşlanmaya dayanıklı kimlik karşılaştırması."),
   ("ByteTrack · yüz takibi",
    "Ardışık karelerde aynı yüzü tek kimlik altında takip eder.",
    "Hareket tahmini + kutu örtüşmesiyle takip; düşük güvenli tespitleri de değerlendirerek kalabalıkta kimlik kaybını azaltır.",
    "Kare 100 ve kare 107 (kesme sonrası) aynı kimliği korur",
    "Kare bazlı tespitleri tek kişinin sürekli görünümüne bağlamak; ekran süresi sinyali."),
   ("HDBSCAN · kümeleme",
    "Aynı videodaki farklı yüz izlerini kişi bazlı gruplara ayırır.",
    "Küme sayısını önceden bilmeden yoğunluk tabanlı kümeleme yapar; gürültüyü ayrı tutar. Az sayıda yüzde bile kararlı sonuç verir.",
    "8 yüz izi → 3 küme: { Sunucu, Konuk A, Konuk B }",
    "Farklı kesimlerdeki aynı yüzü tek kişi altında toplamak."),
   ("pgvector · kişi bankası",
    "Yüz kimlik vektörlerini veritabanında saklayıp hızlı benzerlik araması yapar.",
    "Veritabanına vektör tipi ve benzerlik operatörü ekler. Bir yüz için en yakın kişiler milisaniyeler içinde döner. Vektörler hiçbir arayüzden dışarı verilmez.",
    "Yeni yüz vektörü → bankada en yakın 5 kişi → [ (kişi_001, 0.91), (kişi_002, 0.72) ]",
    "Kimliği kalıcı saklamak; yeni klipleri eski kişilerle eşleştiren büyüyen banka — \"kendini geliştiren\" sistemin belleği."),
  ]),
 ("8 · Sahne / Cisim Tanıma — Planlı (v0.4 hedefi)",
  "Henüz yazılmadı; kütüphaneler kurulu. Buraya TRT'ye özel kültürel veri seti gelecek — camii, köprü, dağ, bayrak, tören gibi bu coğrafyaya ait unsurları tanıyan, kendi verimizle eğitilmiş katman.",
  [
   ("YOLO-World · nesne tespiti (ana motor)",
    "Serbest kelime sözlüğüyle karelerde nesne/kişi/ortam tespit eder (\"kürsü\", \"mikrofon\", \"bayrak\").",
    "Açık-sözlük tanımayı hızlı bir dedektörle birleştirir; eğitimde görmediği kategorileri bile metin tanımından tanıyabilir. Her sahneden birkaç anahtar kare örneklenip sözlükteki sınıflar denenir.",
    "Kare + sözlük [\"kürsü\",\"mikrofon\",\"Türk bayrağı\"] → [ \"Türk bayrağı\" 0.89, kutu ]",
    "\"Bayrak sahnesi\", \"stüdyo paneli\" gibi semantik arşiv aramasını besleyen görsel etiketler."),
   ("SigLIP · atmosfer skorlama",
    "Sahnenin genel atmosferini/ortamını skorlar (\"kapalı stüdyo\", \"açık alan\", \"gece çekimi\"). Logo benzerliğinde de kullanılır.",
    "Görüntü-metin benzerliği modeli. Kareyi aday metin açıklamalarıyla eşleyip her birine benzerlik skoru verir.",
    "Kare + [\"kapalı stüdyo\",\"açık alan\",\"gece\"] → { kapalı stüdyo: 0.91, ... }",
    "Nesne listesinin ötesinde sahnenin bağlamını/atmosferini etiketlemek."),
   ("PySceneDetect · sahne bölme",
    "Videoyu sahnelere böler (kesme noktası tespiti); her sahneye anahtar kare üretir.",
    "Ardışık kareler arası renk/parlaklık farkını ölçer; fark eşiği aşılınca yeni sahne ilan eder.",
    "60 dk haber → yaklaşık 300 sahne segmenti [ (0–12.4sn), (12.4–31.7sn), ... ]",
    "Kare kare yerine sahne bazlı işleme — GPU ve hesap yükünü dramatik düşürür."),
   ("GroundingDINO + RAM (yedek)",
    "Ana motorun lisansı geçmezse yedek nesne dedektörü (GroundingDINO); etiket sözlüğünü genişletmek için keşif aracı (RAM).",
    "GroundingDINO metin tanımından nesne bulur. RAM kapalı sözlük olmadan görüntüdeki her şeyi etiketler, \"arşivde ne tür nesneler geçiyor\" sorusunu veri-odaklı keşfeder.",
    "RAM: kare → [\"kürsü\",\"mikrofon\",\"kravat\",\"Türk bayrağı\",\"stüdyo aydınlatması\"]",
    "Lisans riskine karşı güvence + kültürel etiket sözlüğünü elle yazmadan büyütmek."),
  ]),
]

title_block()
for st, sd, models in DATA:
    sec_title(st)
    sec_desc(sd)
    for name, ne, nasil, ornek, cozer in models:
        model(name)
        field("Ne işe yarar:", ne)
        field("Nasıl çalışır:", nasil)
        ex(ornek)
        field("Çözdüğü sorun:", cozer)

# Elenenler tablosu
sec_title("9 · Denendi ve Elendi")
sec_desc("Test edilen ama seçilmeyen modeller. Bunları denemek, doğru kararı verebilmek için yapıldı.")
rows = [
 ("EasyOCR","OCR adayı","PaddleOCR'a göre net kazanım yok; bağımlılığı OCR ortamına Windows/sürücü riski sokuyor."),
 ("deepseek-ocr","Görsel-dil OCR","Tam başarısız: sayfaların hepsine yazı yerine \"/\" veya komut tekrarı döndürdü."),
 ("LLaVA-NeXT-Video","Video modeliyle jenerik okuma","Kıyaslamada sonuç yetersiz kaldı."),
 ("granite-3.2-8b","Üst denetim modeli","Bağlam tuzaklarına düştü, yanlış kanonik ad üretti."),
 ("glm-4.7-flash","Üst denetim modeli","Çok uzun düşünme, eksik çıktı, normal konuşma ayrımı zayıf."),
 ("qwq-32b","Üst denetim modeli","Format sorunu ve agresif değiştirme riski."),
 ("gpt-oss-20b","Üst denetim modeli","Bu görevde zayıf/orta-alt performans."),
 ("mistral-nemo","Üst denetim modeli","Karar politikası tutarsız."),
 ("Qwen3-14B","Üst denetim modeli","Değiştirme çok agresif; ek düzeltici gerektiriyor."),
 ("qwen3.6:35b (yerel)","Özet motoru (yerel)","Özet formatını tutmadı; soru sorup makale yazdı → Claude'a geçildi."),
 ("qwen2.5-vl (metin)","Üst denetim (metin)","Metin-only testte döngüye girdi. Görsel/bekçi rolünde başarılı (bkz. Bölüm 3)."),
]
tbl = doc.add_table(rows=1, cols=3)
tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
tbl.autofit = False
widths = [Cm(4.0), Cm(4.8), Cm(9.2)]
hdr = tbl.rows[0].cells
for i, h in enumerate(["Model","Ne için denendi","Neden elendi"]):
    hdr[i].width = widths[i]
    pp = hdr[i].paragraphs[0]; pp.paragraph_format.space_after=Pt(0)
    rr = pp.add_run(h); rr.font.bold=True; rr.font.size=Pt(9); rr.font.color.rgb=WHITE; rr.font.name=FONT
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'),'clear'); shd.set(qn('w:fill'),'16315C')
    hdr[i]._tc.get_or_add_tcPr().append(shd)
for r0 in rows:
    cells = tbl.add_row().cells
    for i, val in enumerate(r0):
        cells[i].width = widths[i]
        pp = cells[i].paragraphs[0]; pp.paragraph_format.space_after=Pt(0)
        rr = pp.add_run(val); rr.font.size=Pt(8.6); rr.font.color.rgb=BODY; rr.font.name=FONT
        if i == 0: rr.font.bold = True

note = doc.add_paragraph(); note.paragraph_format.space_before=Pt(8)
rn = note.add_run("Bu belge MITAS mutfak kayıtlarındaki karar günlüğü, güncel durum ve OCR karar dosyalarından, doğrudan koddaki yapılandırmaya bakılarak hazırlanmıştır.")
rn.font.italic=True; rn.font.size=Pt(8.4); rn.font.color.rgb=GREY

out = r"E:\MITAS\mutfak\MITAS_Model_Envanteri.docx"
doc.save(out)
print("OK", out)
